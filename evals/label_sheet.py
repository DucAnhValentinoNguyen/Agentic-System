"""Readable, blind labelling sheet for the judge calibration.

  uv run --with openpyxl python evals/label_sheet.py make    # evals/results/audit_label.xlsx
  uv run --with openpyxl python evals/label_sheet.py merge   # copies the filled column back into audit.csv

The sheet hides the judge's verdict, so the human labels stay independent of it.
"""
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

RES = Path(__file__).resolve().parent / "results"
SHEET = RES / "audit_label.xlsx"
COLS = [("id", 14), ("variant", 8), ("question", 34), ("answer", 70), ("cited_sources", 90), ("human_unsupported", 14)]


def make() -> None:
    rows = list(csv.DictReader(open(RES / "audit.csv", encoding="utf-8")))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "audit"
    ws.append([c for c, _ in COLS])
    for r in rows:
        ws.append([r[c] if c != "human_unsupported" else None for c, _ in COLS])
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, len(COLS)).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "C2"
    wb.save(SHEET)
    print("wrote", SHEET, f"({len(rows)} rows). Fill the yellow column: number of claims in `answer` NOT supported by `cited_sources` (0 if none).")


def merge() -> None:
    ws = openpyxl.load_workbook(SHEET).active
    got = {(str(r[0]), str(r[1])): r[5] for r in ws.iter_rows(min_row=2, values_only=True) if r[5] not in (None, "")}
    rows = list(csv.DictReader(open(RES / "audit.csv", encoding="utf-8")))
    for r in rows:
        r["human_unsupported"] = str(int(got[(r["id"], r["variant"])])) if (r["id"], r["variant"]) in got else ""
    with open(RES / "audit.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print("merged", len(got), "labels into audit.csv; now run: uv run python evals/calibrate.py score")


{"make": make, "merge": merge}[sys.argv[1]]()
